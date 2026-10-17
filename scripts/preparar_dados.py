#!/usr/bin/env python3
"""Extraia exclusivamente campos analíticos sanitizados de uma planilha Porsche.

Uso: python scripts/preparar_dados.py fonte.xlsx [--saida data/vendas-sanitizadas.json]
       [--html index.html]
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


MAX_SAFE_INTEGER = 2**53 - 1
COLUMNS = {
    "id": "sale_id",
    "data_venda": "SaleDateSanitized",
    "modelo": "PorscheModelSanitized",
    "ano_modelo": "ModelYearSanitized",
    "preco_centavos": "SalesPriceSanitized",
    "pagamento": "PayMethodSanitized",
    "cidade": "CitySanitized",
    "estado": "StateSanitized",
    "status_entrega": "DeliveryStatusSanitized",
}
TEXT_FIELDS = ("modelo", "pagamento", "cidade", "estado", "status_entrega")
DATA_TAG = re.compile(
    r'(<script\b(?=[^>]*\bid=["\']vendas-data["\'])(?=[^>]*\btype=["\']application/json["\'])[^>]*>)(.*?)(</script\s*>)',
    re.IGNORECASE | re.DOTALL,
)


class PreparationError(ValueError):
    """Erro que impede produzir um conjunto confiável."""


def parse_price(value: Any) -> int | None:
    """USD decimal completo, não negativo, sem separador de milhar; sem arredondar."""
    if value is None or isinstance(value, bool):
        return None
    if not isinstance(value, (str, int, float, Decimal)):
        return None
    text = str(value)
    if re.fullmatch(r"[0-9]+(?:\.[0-9]{1,2})?", text) is None:
        return None
    # Decimal evita a perda de centavos do ponto flutuante binário.
    cents = int(Decimal(text) * 100)
    if cents > MAX_SAFE_INTEGER:
        raise PreparationError("Preço excede o limite inteiro seguro do JavaScript em centavos.")
    return cents


def parse_model_year(value: Any) -> int | None:
    """Ano inteiro de calendário, sem derivar ou corrigir a partir da data da venda."""
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        return None
    text = str(value)
    if re.fullmatch(r"[0-9]{1,4}", text) is None:
        return None
    year = int(text)
    return year if 1 <= year <= 9999 else None


def parse_sale_date(value: Any) -> str | None:
    """Somente datas ISO reais; INVALID, ausência e datas impossíveis ficam nulas."""
    if not isinstance(value, str) or re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value) is None:
        return None
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        return None


def safe_json(data: Any, *, pretty: bool = True) -> str:
    """JSON seguro inclusive dentro de script HTML, sem sequências de fechamento."""
    result = json.dumps(data, ensure_ascii=False, allow_nan=False, indent=2 if pretty else None)
    return result.replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")


def update_html_text(html: str, data: Any) -> str:
    """Troca apenas o bloco JSON conhecido; preserva todo o script da aplicação."""
    matches = list(DATA_TAG.finditer(html))
    if len(matches) != 1:
        raise PreparationError(
            'O HTML deve conter exatamente um <script id="vendas-data" type="application/json">.'
        )
    match = matches[0]
    return html[: match.start(2)] + "\n" + safe_json(data) + "\n" + html[match.end(2) :]


def prepare(source: Path, *, sheet_name: str = "Sanitized") -> dict[str, Any]:
    """Leia a fonte preservada, valide e gere um objeto sem campos pessoais."""
    source = Path(source)
    if not source.is_file():
        raise PreparationError(f"Planilha não encontrada: {source}")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise PreparationError(f"Aba necessária ausente: {sheet_name}.")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        try:
            header_cells = next(rows)
        except StopIteration as error:
            raise PreparationError("A planilha está vazia.") from error
        headers = [cell.value for cell in header_cells]
        missing = set(COLUMNS.values()) - set(headers)
        if missing:
            raise PreparationError("Colunas necessárias ausentes: " + ", ".join(sorted(missing)))
        duplicates = [name for name, count in Counter(headers).items() if name is not None and count > 1]
        if duplicates:
            raise PreparationError("Cabeçalhos repetidos impedem o mapeamento inequívoco.")
        positions = {field: headers.index(column) for field, column in COLUMNS.items()}
        sales = []
        seen_ids: set[str] = set()
        missing_categories: Counter[str] = Counter()
        blank_rows = 0
        invalid_years = 0
        invalid_dates_marker = 0
        formulas = [cell.coordinate for cell in header_cells if cell.data_type == "f"]
        for row_number, cells in enumerate(rows, start=2):
            formulas.extend(cell.coordinate for cell in cells if cell.data_type == "f")
            if all(cell.value is None for cell in cells):
                blank_rows += 1
                continue
            values = {field: cells[position].value for field, position in positions.items()}
            sale_id = values["id"]
            if isinstance(sale_id, bool) or not isinstance(sale_id, (str, int)) or str(sale_id).strip() == "":
                raise PreparationError(f"ID ausente ou incompatível na linha {row_number}.")
            if isinstance(sale_id, int) and abs(sale_id) > MAX_SAFE_INTEGER:
                raise PreparationError(f"ID numérico excede limite seguro na linha {row_number}.")
            # Trata 6 e '6' como o mesmo identificador, sem descartar nenhum registro.
            id_key = str(sale_id)
            if id_key in seen_ids:
                raise PreparationError(f"ID repetido na linha {row_number}; nenhuma venda foi deduplicada.")
            seen_ids.add(id_key)
            try:
                cents = parse_price(values["preco_centavos"])
            except PreparationError as error:
                raise PreparationError(f"Linha {row_number}: {error}") from error
            year = parse_model_year(values["ano_modelo"])
            invalid_years += year is None
            invalid_dates_marker += values["data_venda"] == "INVALID"
            sale = {
                "id": sale_id,
                "data_venda": parse_sale_date(values["data_venda"]),
                "modelo": values["modelo"],
                "ano_modelo": year,
                "preco_centavos": cents,
                "pagamento": values["pagamento"],
                "cidade": values["cidade"],
                "estado": values["estado"],
                "status_entrega": values["status_entrega"],
            }
            for field in TEXT_FIELDS:
                value = sale[field]
                if value is None or value == "":
                    sale[field] = "Não informado"
                    missing_categories[field] += 1
                elif not isinstance(value, str):
                    raise PreparationError(f"Categoria textual incompatível: {COLUMNS[field]}, linha {row_number}.")
            sales.append(sale)
        if formulas:
            # openpyxl não recalcula. Mesmo havendo cache, não o declaramos validado.
            raise PreparationError(
                f"Encontradas {len(formulas)} fórmula(s) na aba {sheet_name}. "
                "É necessário recalcular e verificar os resultados antes da importação; "
                "este script não considera caches automaticamente confiáveis."
            )
    finally:
        workbook.close()

    commercial = [sale for sale in sales if sale["status_entrega"] != "Cancelled"]
    total_cents = sum(sale["preco_centavos"] for sale in sales if sale["preco_centavos"] is not None)
    if total_cents > MAX_SAFE_INTEGER:
        raise PreparationError("Soma dos preços excede o limite inteiro seguro do JavaScript em centavos.")
    valid_dates = [sale["data_venda"] for sale in sales if sale["data_venda"] is not None]
    metadata = {
        "arquivo": source.name,
        "aba": sheet_name,
        "sha256": source_hash,
        "moeda": "USD",
        "total_importado": len(sales),
        "ids_unicos": len(seen_ids),
        "cancelados": len(sales) - len(commercial),
        "comerciais": len(commercial),
        "datas_invalidas": sum(sale["data_venda"] is None for sale in sales),
        "datas_invalidas_comerciais": sum(sale["data_venda"] is None for sale in commercial),
        "marcadores_INVALID_data": invalid_dates_marker,
        "precos_invalidos": sum(sale["preco_centavos"] is None for sale in sales),
        "precos_invalidos_comerciais": sum(sale["preco_centavos"] is None for sale in commercial),
        "precos_validos": sum(sale["preco_centavos"] is not None for sale in sales),
        "precos_validos_comerciais": sum(sale["preco_centavos"] is not None for sale in commercial),
        "anos_modelo_invalidos": invalid_years,
        "ausencias_categoricas": dict(missing_categories),
        "linhas_vazias_ignoradas": blank_rows,
        "formulas": 0,
        "data_min": min(valid_dates) if valid_dates else None,
        "data_max": max(valid_dates) if valid_dates else None,
        "valor_total_centavos": total_cents,
        "valor_comercial_centavos": sum(sale["preco_centavos"] for sale in commercial if sale["preco_centavos"] is not None),
        "entregues_comerciais": sum(sale["status_entrega"] == "Delivered" for sale in commercial),
    }
    return {"metadata": metadata, "vendas": sales}


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("fonte", type=Path, help="Planilha original preservada e fora do Git.")
    parser.add_argument("--saida", type=Path, default=root / "data/vendas-sanitizadas.json")
    parser.add_argument("--aba", default="Sanitized")
    parser.add_argument("--html", type=Path, help="HTML com o marcador JSON vendas-data.")
    args = parser.parse_args()
    try:
        data = prepare(args.fonte, sheet_name=args.aba)
        # Valide o HTML antes de escrever qualquer saída.
        updated_html = update_html_text(args.html.read_text(encoding="utf-8"), data) if args.html else None
        args.saida.parent.mkdir(parents=True, exist_ok=True)
        args.saida.write_text(safe_json(data) + "\n", encoding="utf-8")
        if args.html is not None:
            args.html.write_text(updated_html, encoding="utf-8")
    except (PreparationError, OSError) as error:
        parser.exit(1, f"Falha na preparação: {error}\n")
    print(json.dumps(data["metadata"], ensure_ascii=False, indent=2))
    print(f"JSON gerado: {args.saida}")
    if args.html:
        print(f"Dados incorporados atualizados: {args.html}")


if __name__ == "__main__":
    main()
