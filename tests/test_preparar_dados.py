"""Testes sem rede: conversão, falhas relevantes e reconciliação da base publicada."""

from collections import Counter
from decimal import Decimal
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("preparar_dados", ROOT / "scripts/preparar_dados.py")
prep = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(prep)


class ConversionTests(unittest.TestCase):
    def test_decimal_examples_and_zero(self):
        self.assertEqual(prep.parse_price("79500.00"), 7950000)
        self.assertEqual(prep.parse_price("153200.50"), 15320050)
        self.assertEqual(prep.parse_price("0.00"), 0)
        self.assertEqual(prep.parse_price("0.01"), 1)
        self.assertEqual(prep.parse_price(Decimal("0.29")), 29)

    def test_missing_invalid_and_strict_prices(self):
        for value in (None, "", "INVALID", "USD 79.00", "1,000.00", "79abc", "1.001", "-1", "NaN", "Infinity", "1e3", " 1.00 ", True, float("inf"), float("nan")):
            with self.subTest(value=value):
                self.assertIsNone(prep.parse_price(value))

    def test_safe_integer_boundary(self):
        self.assertEqual(prep.parse_price("90071992547409.91"), prep.MAX_SAFE_INTEGER)
        with self.assertRaisesRegex(prep.PreparationError, "limite inteiro seguro"):
            prep.parse_price("90071992547409.92")

    def test_model_year(self):
        self.assertEqual(prep.parse_model_year("2024"), 2024)
        for value in (None, "INVALID", "2024.0", 2024.0, True, "0", "10000", "-1"):
            with self.subTest(value=value):
                self.assertIsNone(prep.parse_model_year(value))

    def test_calendar_dates_and_future_dates(self):
        self.assertEqual(prep.parse_sale_date("2024-03-14"), "2024-03-14")
        self.assertEqual(prep.parse_sale_date("2024-02-29"), "2024-02-29")
        self.assertEqual(prep.parse_sale_date("2027-12-25"), "2027-12-25")
        for value in ("INVALID", "2024-02-30", "2023-02-29", "2024-13-01", "2024-01-00", "20240314", "2024-3-14", None, "2024-03-14T00:00:00"):
            with self.subTest(value=value):
                self.assertIsNone(prep.parse_sale_date(value))

    def test_safe_embedding_and_application_preservation(self):
        payload = {"vendas": [{"modelo": "</script><script>alert(1)</script>\u2028\u2029"}]}
        before = '<html><script id="vendas-data" type="application/json">{}</script><script>const app = "preservado";</script></html>'
        after = prep.update_html_text(before, payload)
        self.assertIn('<script>const app = "preservado";</script>', after)
        embedded = prep.DATA_TAG.search(after).group(2)
        self.assertNotIn("<", embedded)
        self.assertNotIn("\u2028", embedded)
        self.assertNotIn("\u2029", embedded)
        self.assertEqual(json.loads(embedded), payload)

    def test_missing_or_duplicate_html_marker_fails(self):
        for html in ("<html></html>", '<script id="vendas-data" type="text/javascript">{}</script>', '<script id="vendas-data" type="application/json">{}</script>' * 2):
            with self.subTest(html=html), self.assertRaises(prep.PreparationError):
                prep.update_html_text(html, {})


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "fixture.xlsx"

    def write_book(self, patches):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sanitized"
        sheet.append(list(prep.COLUMNS.values()) + ["customer_name", "salesperson"])
        base = {
            "id": 6, "data_venda": "2024-03-14", "modelo": "911 GT3",
            "ano_modelo": "2024", "preco_centavos": "153200.50", "pagamento": "Wire Transfer",
            "cidade": "Boston", "estado": "MA", "status_entrega": "Delivered",
        }
        for patch in patches:
            row = base | patch
            sheet.append([row[field] for field in prep.COLUMNS] + ["DADO PESSOAL TESTE", "NOME TESTE"])
        workbook.save(self.path)
        workbook.close()

    def test_invalid_values_preserved_and_coverage_reported(self):
        self.write_book([
            {"data_venda": "INVALID", "preco_centavos": None, "ano_modelo": "INVALID"},
            {"id": 7, "preco_centavos": "0.00", "pagamento": "Bank Transfer"},
            {"id": 8, "data_venda": "2024-02-30", "status_entrega": "Cancelled", "pagamento": "ACH Payment"},
        ])
        data = prep.prepare(self.path)
        metadata = data["metadata"]
        self.assertEqual(metadata["total_importado"], 3)
        self.assertEqual(metadata["cancelados"], 1)
        self.assertEqual(metadata["datas_invalidas"], 2)
        self.assertEqual(metadata["datas_invalidas_comerciais"], 1)
        self.assertEqual(metadata["anos_modelo_invalidos"], 1)
        self.assertEqual(metadata["precos_invalidos_comerciais"], 1)
        self.assertEqual(metadata["precos_validos_comerciais"], 1)
        self.assertEqual(metadata["valor_comercial_centavos"], 0)
        self.assertEqual(metadata["valor_total_centavos"], 15320050)
        self.assertIsNone(data["vendas"][0]["preco_centavos"])
        self.assertEqual(data["vendas"][1]["preco_centavos"], 0)
        self.assertEqual({s["pagamento"] for s in data["vendas"]}, {"Wire Transfer", "Bank Transfer", "ACH Payment"})
        self.assertNotIn("DADO PESSOAL", prep.safe_json(data))
        self.assertEqual(set(data["vendas"][0]), set(prep.COLUMNS))

    def test_duplicate_id_fails_without_deduplication(self):
        self.write_book([{}, {"id": "6"}])
        with self.assertRaisesRegex(prep.PreparationError, "ID repetido"):
            prep.prepare(self.path)

    def test_identical_model_city_price_does_not_deduplicate(self):
        self.write_book([{}, {"id": 7}])
        self.assertEqual(prep.prepare(self.path)["metadata"]["total_importado"], 2)

    def test_formula_fails_instead_of_assuming_cache_zero(self):
        self.write_book([{"preco_centavos": "=100+200"}])
        with self.assertRaisesRegex(prep.PreparationError, "fórmula"):
            prep.prepare(self.path)

    def test_aggregate_overflow_fails(self):
        self.write_book([{"preco_centavos": "90071992547409.91"}, {"id": 7, "preco_centavos": "0.01"}])
        with self.assertRaisesRegex(prep.PreparationError, "Soma dos preços excede"):
            prep.prepare(self.path)

    def test_regeneration_is_reproducible_and_source_unchanged(self):
        self.write_book([{}, {"id": 7}])
        before = self.path.read_bytes()
        first = prep.prepare(self.path)
        second = prep.prepare(self.path)
        self.assertEqual(prep.safe_json(first), prep.safe_json(second))
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(first["metadata"]["sha256"], hashlib.sha256(before).hexdigest())

    def test_missing_required_column_fails(self):
        workbook = Workbook()
        workbook.active.title = "Sanitized"
        workbook.active.append(["sale_id"])
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(prep.PreparationError, "Colunas necessárias ausentes"):
            prep.prepare(self.path)


class PublishedDataTests(unittest.TestCase):
    def test_published_json_reconciles_and_has_no_personal_fields(self):
        data = json.loads((ROOT / "data/vendas-sanitizadas.json").read_text(encoding="utf-8"))
        sales, metadata = data["vendas"], data["metadata"]
        commercial = [sale for sale in sales if sale["status_entrega"] != "Cancelled"]
        self.assertEqual(len({str(sale["id"]) for sale in sales}), len(sales))
        self.assertEqual(len(sales), metadata["total_importado"])
        self.assertEqual(len(commercial), metadata["comerciais"])
        self.assertEqual(sum(s["preco_centavos"] or 0 for s in sales), metadata["valor_total_centavos"])
        self.assertEqual(sum(s["preco_centavos"] or 0 for s in commercial), metadata["valor_comercial_centavos"])
        self.assertEqual(sum(s["data_venda"] is None for s in commercial), metadata["datas_invalidas_comerciais"])
        for sale in sales:
            self.assertEqual(set(sale), set(prep.COLUMNS))
        for field in ("modelo", "estado", "pagamento"):
            quantities = Counter(sale[field] for sale in commercial)
            grouped_cents = {}
            for sale in commercial:
                grouped_cents[sale[field]] = grouped_cents.get(sale[field], 0) + (sale["preco_centavos"] or 0)
            self.assertEqual(sum(quantities.values()), len(commercial))
            self.assertEqual(sum(grouped_cents.values()), metadata["valor_comercial_centavos"])

    @unittest.skipUnless(os.environ.get("PORSCHE_SOURCE"), "Defina PORSCHE_SOURCE para conferir a fonte original preservada.")
    def test_real_source_independent_decimal_audit_and_reproducibility(self):
        source = Path(os.environ["PORSCHE_SOURCE"])
        published = json.loads((ROOT / "data/vendas-sanitizadas.json").read_text(encoding="utf-8"))
        self.assertEqual(prep.prepare(source), published)
        # Segunda leitura e cálculo com Decimal, sem chamar conversores da preparação.
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            rows = workbook["Sanitized"].iter_rows(values_only=True)
            headers = next(rows)
            records = [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]
        finally:
            workbook.close()
        commercial = [r for r in records if r["DeliveryStatusSanitized"] != "Cancelled"]
        total = sum((Decimal(r["SalesPriceSanitized"]) for r in records), Decimal(0))
        total_commercial = sum((Decimal(r["SalesPriceSanitized"]) for r in commercial), Decimal(0))
        self.assertEqual(int(total * 100), published["metadata"]["valor_total_centavos"])
        self.assertEqual(int(total_commercial * 100), published["metadata"]["valor_comercial_centavos"])
        self.assertEqual(len(records), published["metadata"]["total_importado"])
        self.assertEqual(len(commercial), published["metadata"]["comerciais"])
        for source_field, output_field in (("PorscheModelSanitized", "modelo"), ("StateSanitized", "estado"), ("PayMethodSanitized", "pagamento")):
            reference_groups = {}
            for row in commercial:
                key = row[source_field]
                count, amount = reference_groups.get(key, (0, Decimal(0)))
                reference_groups[key] = (count + 1, amount + Decimal(row["SalesPriceSanitized"]))
            for key, (count, amount) in reference_groups.items():
                group = [sale for sale in published["vendas"] if sale["status_entrega"] != "Cancelled" and sale[output_field] == key]
                self.assertEqual(len(group), count)
                self.assertEqual(sum(sale["preco_centavos"] for sale in group), int(amount * 100))


if __name__ == "__main__":
    unittest.main()
