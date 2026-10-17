"""Browser checks against an independent XLSX oracle; no web server is needed.

Run: PORSCHE_SOURCE=/path/to/source.xlsx python -m unittest discover -s tests -p 'test_dashboard_browser.py' -v
Requires openpyxl, playwright and Chromium (CHROMIUM_PATH overrides its location).
Behavior checks load the local HTML using set_content while the browser is offline.
Direct file:// navigation is checked separately and explicitly skipped if a managed
browser policy blocks local files. No browser policy is changed.
"""

import copy
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP
import json
import os
from pathlib import Path
import re
import tempfile
import unittest

from openpyxl import load_workbook
from playwright.sync_api import Error as PlaywrightError, sync_playwright


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(os.environ["PORSCHE_SOURCE"]).resolve() if os.environ.get("PORSCHE_SOURCE") else None


def source_records(path):
    """Read the workbook directly, without importing the preparation script."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = workbook["Sanitized"].iter_rows(values_only=True)
        headers = next(rows)
        result = []
        for cells in rows:
            if all(cell is None for cell in cells):
                continue
            values = dict(zip(headers, cells))
            if any(isinstance(cell, str) and cell.startswith("=") for cell in cells):
                raise AssertionError("Formula found: cached values need independent verification")
            raw_date = values["SaleDateSanitized"]
            try:
                date = dt.date.fromisoformat(str(raw_date)).isoformat()
            except (TypeError, ValueError):
                date = None
            raw_price = values["SalesPriceSanitized"]
            price = None
            if re.fullmatch(r"\d+(?:\.\d{1,2})?", str(raw_price)):
                price = int(Decimal(str(raw_price)) * 100)
            result.append({
                "id": int(values["sale_id"]),
                "data_venda": date,
                "modelo": values["PorscheModelSanitized"],
                "ano_modelo": int(values["ModelYearSanitized"]),
                "preco_centavos": price,
                "pagamento": values["PayMethodSanitized"],
                "cidade": values["CitySanitized"],
                "estado": values["StateSanitized"],
                "status_entrega": values["DeliveryStatusSanitized"],
            })
        return result
    finally:
        workbook.close()


def expected_summary(rows):
    priced = [row["preco_centavos"] for row in rows if row["preco_centavos"] is not None]
    return {
        "count": len(rows),
        "totalCents": sum(priced) if priced or not rows else None,
        "pricedCount": len(priced),
        "ticketCents": sum(priced) / len(priced) if priced else None,
        "delivered": sum(row["status_entrega"] == "Delivered" for row in rows),
        "invalidDates": sum(row["data_venda"] is None for row in rows),
    }


def cents_in_text(text):
    """USD display uses exactly two fractional digits; punctuation is locale independent."""
    digits = re.sub(r"\D", "", text)
    return int(digits) if digits else None


def usd_number(cents):
    amount = (Decimal(cents) / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{amount:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


class DashboardBrowserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if SOURCE is None or not SOURCE.is_file():
            raise RuntimeError("Set PORSCHE_SOURCE to the original .xlsx file for the independent oracle")
        cls.imported = source_records(SOURCE)
        cls.rows = [row for row in cls.imported if row["status_entrega"] != "Cancelled"]
        cls.playwright = sync_playwright().start()
        cls.browser = cls.playwright.chromium.launch(
            executable_path=os.environ.get("CHROMIUM_PATH", "/usr/bin/chromium"),
            headless=True,
        )
        cls.workdir = tempfile.TemporaryDirectory(prefix="porsche-browser-")

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.playwright.stop()
        cls.workdir.cleanup()

    def setUp(self):
        self.context = self.browser.new_context(offline=True, viewport={"width": 1440, "height": 1000})
        self.page = self.context.new_page()
        self.errors = []
        self.external_requests = []
        self.page.on("pageerror", lambda error: self.errors.append(str(error)))
        self.page.on("console", lambda message: self.errors.append(message.text) if message.type == "error" else None)
        self.page.on("request", lambda request: self.external_requests.append(request.url)
                     if request.url.startswith(("https://", "http://")) else None)
        self.page.set_content((ROOT / "index.html").read_text(encoding="utf-8"))
        self.page.wait_for_function("Boolean(window.PorscheDashboard)")

    def tearDown(self):
        try:
            self.assertEqual(self.errors, [], "JavaScript or console errors")
            self.assertEqual(self.external_requests, [], "Self-contained HTML attempted a network request")
        finally:
            self.context.close()

    def clear(self):
        self.page.locator("#clear-filters").click()

    def choose_city(self, row):
        options = self.page.locator("#filter-cidade option").evaluate_all(
            "options => options.map(o => ({value:o.value, text:o.textContent}))"
        )
        options = [item for item in options if row["cidade"] in item["text"]]
        matching = [item for item in options if re.search(r"\b" + re.escape(row["estado"]) + r"\b", item["text"])]
        self.assertEqual(len(matching or options), 1, "City option must identify an unambiguous city/state")
        self.page.select_option("#filter-cidade", (matching or options)[0]["value"])

    def assert_groups(self, rows):
        groups = self.page.evaluate("PorscheDashboard.getGroups()")
        whole = expected_summary(rows)
        for api_key, dom_key, field in (("modelos", "models", "modelo"), ("estados", "states", "estado"), ("pagamentos", "payments", "pagamento")):
            with self.subTest(groups=api_key):
                grouped = {}
                for row in rows:
                    grouped.setdefault(row[field], []).append(row)
                actual = groups[api_key]
                self.assertEqual({group["key"] for group in actual}, set(grouped))
                self.assertEqual(sum(group["count"] for group in actual), whole["count"])
                if whole["totalCents"] is not None:
                    self.assertEqual(sum(group["totalCents"] or 0 for group in actual), whole["totalCents"])
                self.assertAlmostEqual(sum(group["share"] for group in actual), 1 if rows else 0)
                for group in actual:
                    expected = expected_summary(grouped[group["key"]])
                    for key in ("count", "totalCents", "pricedCount", "delivered", "invalidDates"):
                        self.assertEqual(group[key], expected[key], (group["key"], key))
                    if expected["ticketCents"] is None:
                        self.assertIsNone(group["ticketCents"])
                    else:
                        self.assertAlmostEqual(group["ticketCents"], expected["ticketCents"], delta=0.51)
                    self.assertAlmostEqual(group["share"], len(grouped[group["key"]]) / len(rows))
                    expected_value_share = group["totalCents"] / whole["totalCents"] if whole["totalCents"] and group["totalCents"] is not None else None
                    if expected_value_share is None:
                        self.assertIsNone(group["valueShare"])
                    else:
                        self.assertAlmostEqual(group["valueShare"], expected_value_share)
                measure = "totalCents" if field == "modelo" else "count"
                values = [group[measure] if group[measure] is not None else -1 for group in actual]
                self.assertEqual(values, sorted(values, reverse=True))
                chart = self.page.locator(f"#{dom_key}-chart .bar-row").evaluate_all("""rows => rows.map(row => ({
                    key: row.dataset.category, count: Number(row.dataset.count),
                    totalCents: row.dataset.totalCents === '' ? null : Number(row.dataset.totalCents)
                }))""")
                self.assertEqual(chart, [{key: group[key] for key in ("key", "count", "totalCents")} for group in actual[:len(chart)]])
                self.assertIn(len(chart), {min(10, len(actual)), len(actual)})
                support = self.page.locator(f"#{dom_key}-support tr").evaluate_all(
                    "rows => rows.map(row => Array.from(row.children, cell => cell.textContent))"
                )
                if rows:
                    self.assertEqual(len(support), len(actual))
                    for cells, group in zip(support, actual):
                        self.assertIn(group["key"], cells[0])
                        self.assertEqual(cents_in_text(cells[1]), group["count"])
                        if field == "modelo":
                            self.assertEqual(cents_in_text(cells[2]), group["totalCents"])
                            ticket = group["ticketCents"]
                            expected_ticket = int(Decimal(str(ticket)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if ticket is not None else None
                            self.assertEqual(cents_in_text(cells[3]), expected_ticket)
                        else:
                            displayed_share = float(cells[2].replace("%", "").replace(",", ".").strip())
                            self.assertAlmostEqual(displayed_share, group["share"] * 100, delta=0.051)
                            self.assertEqual(cents_in_text(cells[3]), group["totalCents"])
                conclusion = self.page.locator(f"#{dom_key}-conclusion").inner_text()
                if not rows:
                    self.assertIn("Nenhuma venda", conclusion)
                elif field == "modelo" and whole["totalCents"] is None:
                    self.assertIn("indisponíveis", conclusion)
                else:
                    leader = actual[0]
                    leaders = [group for group in actual if group[measure] == leader[measure]]
                    for group in leaders[:3]:
                        self.assertIn(group["key"], conclusion)
                    if len(leaders) > 1:
                        self.assertIn("Empate", conclusion)
                    if field == "modelo":
                        self.assertIn(usd_number(leader["totalCents"]), conclusion)
                    else:
                        self.assertIn(str(leader["count"]) + " venda", conclusion)
                        displayed_share = (Decimal(leader["count"]) * 100 / len(rows)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                        self.assertIn(str(displayed_share).replace(".", ",") + "%", conclusion)

    def open_fixture(self, rows, name):
        """Replace the embedded JSON only in an untracked, temporary HTML copy."""
        html = (ROOT / "index.html").read_text(encoding="utf-8")
        pattern = re.compile(r'(<script\b[^>]*\bid="vendas-data"[^>]*>)(.*?)(</script>)', re.S)
        match = pattern.search(html)
        self.assertIsNotNone(match)
        payload = json.loads(match.group(2))
        payload["vendas"] = rows
        payload["metadata"].update({"total_importado": len(rows), "ids_unicos": len(rows), "cancelados": 0, "comerciais": len(rows)})
        content = json.dumps(payload, ensure_ascii=False).replace("<", "\\u003c")
        fixture = Path(self.workdir.name) / (name + ".html")
        fixture.write_text(html[:match.start(2)] + content + html[match.end(2):], encoding="utf-8")
        self.page.goto("about:blank")
        self.page.set_content(fixture.read_text(encoding="utf-8"))
        self.page.wait_for_function("Boolean(window.PorscheDashboard)")

    def assert_view(self, expected):
        self.page.wait_for_timeout(30)
        ids = sorted(row["id"] for row in expected)
        selected = self.page.evaluate("PorscheDashboard.getSelected()")
        self.assertEqual(sorted(row["id"] for row in selected), ids)
        actual = self.page.evaluate("PorscheDashboard.getSummary()")
        oracle = expected_summary(expected)
        for key in ("count", "totalCents", "pricedCount", "delivered", "invalidDates"):
            self.assertEqual(actual[key], oracle[key], key)
        if oracle["ticketCents"] is None:
            self.assertIsNone(actual["ticketCents"])
        else:
            self.assertAlmostEqual(actual["ticketCents"], oracle["ticketCents"], delta=0.51)
        self.assertEqual(cents_in_text(self.page.locator("#kpi-count").inner_text()), len(expected))
        self.assertEqual(cents_in_text(self.page.locator("#kpi-delivered").inner_text()), oracle["delivered"])
        if expected and not oracle["pricedCount"]:
            self.assertIsNone(cents_in_text(self.page.locator("#kpi-value").inner_text()))
        else:
            self.assertEqual(cents_in_text(self.page.locator("#kpi-value").inner_text()), oracle["totalCents"])
        ticket_text = self.page.locator("#kpi-ticket").inner_text()
        if oracle["pricedCount"]:
            expected_ticket = int((Decimal(oracle["totalCents"]) / oracle["pricedCount"]).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            self.assertEqual(cents_in_text(ticket_text), expected_ticket)
        else:
            self.assertIsNone(cents_in_text(ticket_text))
        table_rows = self.page.locator("#sales-body tr[data-sale-id]")
        table = table_rows.evaluate_all("""rows => rows.map(row => ({
            id: Number(row.dataset.saleId), text: row.textContent,
            cells: Array.from(row.querySelectorAll('td'), cell => cell.textContent)
        }))""")
        self.assertEqual(sorted(row["id"] for row in table), ids)
        by_id = {row["id"]: row for row in table}
        for row in expected:
            table_row = by_id[row["id"]]
            text = table_row["text"]
            for key in ("modelo", "cidade", "estado", "pagamento"):
                self.assertIn(str(row[key]), text)
            self.assertIn(str(row["ano_modelo"]), text)
            status_text = {
                "Delivered": ("Delivered", "Entregue"),
                "In Transit": ("In Transit", "Em trânsito"),
                "Pending": ("Pending", "Pendente"),
                "Shipped": ("Shipped", "Enviado"),
                "Awaiting Delivery": ("Awaiting Delivery", "Aguardando entrega"),
                "Awaiting Pickup": ("Awaiting Pickup", "Aguardando retirada"),
                "Awaiting Review": ("Awaiting Review", "Aguardando revisão"),
                "Pending Approval": ("Pending Approval", "Aprovação pendente"),
                "Pending Review": ("Pending Review", "Revisão pendente"),
            }.get(row["status_entrega"], (row["status_entrega"],))
            self.assertTrue(any(label.casefold() in text.casefold() for label in status_text), row["status_entrega"])
            cells = table_row["cells"]
            self.assertEqual(len(cells), 8)
            date_text = cells[4]
            if row["data_venda"]:
                date = dt.date.fromisoformat(row["data_venda"])
                self.assertTrue(row["data_venda"] in date_text or date.strftime("%d/%m/%Y") in date_text, date_text)
            else:
                self.assertRegex(date_text.casefold(), r"sem data|inválid|indisponível|—|–")
            self.assertEqual(cents_in_text(cells[7]), row["preco_centavos"])
        self.assert_groups(expected)

    def test_01_source_and_initial_view_offline(self):
        self.assertEqual(len(self.imported), len({row["id"] for row in self.imported}), "Duplicate source IDs")
        self.assert_view(self.rows)
        self.assertTrue(any(row["data_venda"] is None for row in self.rows))
        self.assertFalse(any(row["status_entrega"] == "Cancelled" for row in self.page.evaluate("PorscheDashboard.getSelected()")))
        for name, field in (("modelo", "modelo"), ("estado", "estado"), ("ano-modelo", "ano_modelo"), ("pagamento", "pagamento")):
            values = self.page.locator(f"#filter-{name} option").evaluate_all("options => options.map(o => o.value)")
            self.assertEqual(values[0], "")
            self.assertEqual(set(values[1:]), {str(row[field]) for row in self.rows})
            if name == "ano-modelo":
                self.assertEqual([int(value) for value in values[1:]], sorted(int(value) for value in values[1:]))
        years = self.page.locator("#filter-ano-venda option").evaluate_all("options => options.map(o => o.value)")
        self.assertEqual(set(years) - {"", "__sem_data__"}, {row["data_venda"][:4] for row in self.rows if row["data_venda"]})
        numeric_years = [int(value) for value in years if value not in ("", "__sem_data__")]
        self.assertEqual(numeric_years, sorted(numeric_years))
        city_values = self.page.locator("#filter-cidade option").evaluate_all("options => options.map(o => o.value)")
        self.assertEqual(len(city_values) - 1, len({(row["cidade"], row["estado"]) for row in self.rows}))

    def test_02_each_filter_independently(self):
        sample = next(row for row in self.rows if row["data_venda"])
        cases = [
            ("modelo", "modelo", sample["modelo"]),
            ("estado", "estado", sample["estado"]),
            ("ano-modelo", "ano_modelo", str(sample["ano_modelo"])),
            ("pagamento", "pagamento", sample["pagamento"]),
        ]
        for selector, key, value in cases:
            with self.subTest(filter=selector, value=value):
                self.clear()
                self.page.select_option("#filter-" + selector, value)
                self.assert_view([row for row in self.rows if str(row[key]) == value])
        self.clear()
        self.choose_city(sample)
        self.assert_view([row for row in self.rows if (row["cidade"], row["estado"]) == (sample["cidade"], sample["estado"])])

    def test_03_sale_year_and_invalid_dates(self):
        self.page.select_option("#filter-ano-venda", "__sem_data__")
        self.assert_view([row for row in self.rows if row["data_venda"] is None])
        self.clear()
        self.page.select_option("#filter-ano-venda", "2024")
        selected = [row for row in self.rows if row["data_venda"] and row["data_venda"].startswith("2024-")]
        self.assert_view(selected)
        self.assertEqual(self.page.evaluate("PorscheDashboard.getSummary().invalidDates"), 0)
        self.clear()
        self.assert_view(self.rows)

    def test_04_combination_empty_result_and_clear(self):
        sample = self.rows[0]
        self.page.select_option("#filter-modelo", sample["modelo"])
        self.page.select_option("#filter-estado", sample["estado"])
        self.assert_view([row for row in self.rows if row["modelo"] == sample["modelo"] and row["estado"] == sample["estado"]])
        self.clear()
        models = sorted({row["modelo"] for row in self.rows})
        states = sorted({row["estado"] for row in self.rows})
        absent = next(((model, state) for model in models for state in states
                       if not any(row["modelo"] == model and row["estado"] == state for row in self.rows)), None)
        self.assertIsNotNone(absent, "This source should provide a real empty combination")
        self.page.select_option("#filter-modelo", absent[0])
        self.page.select_option("#filter-estado", absent[1])
        self.assert_view([])
        self.assertRegex(self.page.locator("#sales-body").inner_text().lower(), r"nenhum|sem resultado|não há|não encontrad")
        self.clear()
        self.assert_view(self.rows)
        for selector in ("modelo", "estado", "cidade", "ano-modelo", "ano-venda", "pagamento"):
            self.assertEqual(self.page.locator("#filter-" + selector).input_value(), "")

    def test_05_keyboard_and_labels(self):
        for name in ("modelo", "estado", "cidade", "ano-modelo", "ano-venda", "pagamento"):
            self.assertGreater(self.page.locator(f'label[for="filter-{name}"]').count(), 0)
        self.page.locator("#filter-modelo").focus()
        self.page.keyboard.press("ArrowDown")
        self.page.keyboard.press("Enter")
        value = self.page.locator("#filter-modelo").input_value()
        self.assertNotEqual(value, "")
        self.assert_view([row for row in self.rows if row["modelo"] == value])
        self.page.locator("#clear-filters").focus()
        self.page.keyboard.press("Enter")
        self.assert_view(self.rows)
        focused = self.page.evaluate("""() => {
            const style = getComputedStyle(document.activeElement);
            return {id: document.activeElement.id, outline: style.outlineStyle,
                    width: parseFloat(style.outlineWidth), shadow: style.boxShadow};
        }""")
        self.assertEqual(focused["id"], "clear-filters")
        self.assertTrue((focused["outline"] != "none" and focused["width"] > 0) or focused["shadow"] != "none", "Focused action needs a visible focus indicator")

    def test_06_responsive_layout_and_no_document_overflow(self):
        for width in (360, 1440):
            with self.subTest(width=width):
                self.page.set_viewport_size({"width": width, "height": 1000})
                self.assert_view(self.rows)
                dimensions = self.page.evaluate("""() => ({
                    scroll: document.documentElement.scrollWidth,
                    client: document.documentElement.clientWidth,
                    body: document.body.scrollWidth
                })""")
                self.assertLessEqual(dimensions["scroll"], dimensions["client"] + 1, dimensions)
                self.assertLessEqual(dimensions["body"], dimensions["client"] + 1, dimensions)
                for selector in ("#filter-modelo", "#filter-estado", "#filter-cidade", "#filter-ano-modelo", "#filter-ano-venda", "#filter-pagamento", "#kpi-value"):
                    box = self.page.locator(selector).bounding_box()
                    self.assertGreater(box["width"], 0)
                    self.assertGreaterEqual(box["x"], -1)
                    self.assertLessEqual(box["x"] + box["width"], width + 1)

    def test_07_rankings_can_expand_to_all_groups(self):
        groups = self.page.evaluate("PorscheDashboard.getGroups()")
        for api_key, kind in (("modelos", "models"), ("estados", "states"), ("pagamentos", "payments")):
            button = self.page.locator(f"#{kind}-expand")
            if len(groups[api_key]) > 10:
                self.assertEqual(self.page.locator(f"#{kind}-chart .bar-row").count(), 10)
                button.click()
                self.assertEqual(button.get_attribute("aria-expanded"), "true")
                self.assertEqual(self.page.locator(f"#{kind}-chart .bar-row").count(), len(groups[api_key]))
                self.assert_groups(self.rows)
                button.click()
                self.assertEqual(self.page.locator(f"#{kind}-chart .bar-row").count(), 10)
            else:
                self.assertTrue(button.is_hidden())

    def test_08_empty_missing_and_zero_price_fixtures(self):
        zero = copy.deepcopy(self.rows[0])
        zero["preco_centavos"] = 0
        missing = copy.deepcopy(self.rows[1])
        missing["preco_centavos"] = None
        for name, fixture in (("empty", []), ("zero", [zero]), ("missing", [missing]), ("partial", [zero, missing])):
            with self.subTest(fixture=name):
                self.open_fixture(fixture, name)
                self.assert_view(fixture)

    def test_09_direct_file_navigation(self):
        try:
            self.page.goto((ROOT / "index.html").as_uri())
        except PlaywrightError as error:
            if "net::ERR_BLOCKED_BY_ADMINISTRATOR" in str(error):
                self.skipTest("file:// navigation blocked by the environment's managed browser policy; offline set_content behavior is tested separately")
            raise
        self.page.wait_for_function("Boolean(window.PorscheDashboard)")
        self.assert_view(self.rows)


if __name__ == "__main__":
    unittest.main(verbosity=2)
